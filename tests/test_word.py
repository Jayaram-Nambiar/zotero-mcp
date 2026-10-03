"""Word field tests. They build a .docx and do not call Zotero."""

from __future__ import annotations

import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock
from xml.etree import ElementTree

from docx import Document

from zotero_mcp import client, word
from zotero_mcp.server import embed_zotero_word_fields
from zotero_mcp.word import (
    EmbeddedItem,
    _bibliography_text,
    _Job,
    _visible_citation,
    chunk_property,
    citation_instruction,
    citation_keys,
    document_preferences,
    embed_document,
    parse_marker_body,
    style_id,
    style_name,
)

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_CUSTOM = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"


def _sample_item(title: str = "Example & <Trial>", family: str = "Smith", year: int = 2020) -> EmbeddedItem:
    return EmbeddedItem(
        csl={
            "type": "article-journal",
            "title": title,
            "author": [{"family": family, "given": "Ada"}],
            "issued": {"date-parts": [[year]]},
        },
        bibliography=f"{family} A. {title}. Journal. {year}.",
    )


def _job(items: dict[str, EmbeddedItem], *, numeric: bool) -> _Job:
    return _Job(
        user_id="1",
        style_id="http://www.zotero.org/styles/apa",
        locale="en-US",
        items=items,
        order={key: index for index, key in enumerate(items, start=1)},
        numeric=numeric,
        session_id="session",
    )


def _write_docx(path: Path, *paragraphs: str) -> Path:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)
    return path


_BODY_KEYS = [f"BODY{index:04d}" for index in range(40)]
_CELL_KEYS = [f"CELL{index:04d}" for index in range(50)]
_MANY_KEYS = _BODY_KEYS[:20] + _CELL_KEYS + _BODY_KEYS[20:]


def _write_many_markers(path: Path) -> None:
    """Save 20 marker paragraphs, a 10 x 5 table of markers, then 20 more paragraphs.

    Dozens of markers, because a paragraph walk that remembers bare id()s starts
    skipping only once earlier element proxies are freed and their ids reused.
    """
    document = Document()
    for key in _BODY_KEYS[:20]:
        document.add_paragraph("Claim {{zotero:" + key + "}}.")
    table = document.add_table(rows=10, cols=5)
    for index, key in enumerate(_CELL_KEYS):
        table.cell(index // 5, index % 5).text = "{{zotero:" + key + "}}"
    for key in _BODY_KEYS[20:]:
        document.add_paragraph("Claim {{zotero:" + key + "}}.")
    document.save(path)


class WordFieldTests(unittest.TestCase):
    def setUp(self) -> None:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)

    def _embed(self, source: Path, items: dict[str, EmbeddedItem], style: str = "vancouver") -> Path:
        destination = self.root / "out.docx"
        embed_document(source, destination, user_id="1", style=style, locale="en-US", items=items, insert_bibliography=True)
        return destination

    def test_marker_options(self) -> None:
        marker = parse_marker_body("ABCD1234+EFGH5678|locator=12|label=page|prefix=see")
        self.assertEqual(marker.kind, "citation")
        assert marker.citation is not None
        self.assertEqual(marker.citation.keys, ["ABCD1234", "EFGH5678"])
        self.assertEqual(marker.citation.locator, "12")
        self.assertEqual(marker.citation.prefix, "see")
        with self.assertRaises(ValueError):
            parse_marker_body("not-a-key")

    def test_marker_errors_quote_the_marker(self) -> None:
        source = _write_docx(self.root / "draft.docx", "Fine {{zotero:ABCD1234}}.", "Typo {{zotero:ABC123}}.")
        with self.assertRaises(ValueError) as raised:
            citation_keys(source)
        self.assertIn("{{zotero:ABC123}}", str(raised.exception))

    def test_preferences_chunk_at_255_characters(self) -> None:
        chunks = chunk_property("a" * 256)
        self.assertEqual(chunks, ["a" * 255, "a"])
        raw = document_preferences(style="vancouver", locale="en-US", session_id="abc", has_bibliography=True)
        payload = json.loads("".join(chunk_property(raw)))
        self.assertEqual(payload["dataVersion"], 4)
        self.assertEqual(payload["prefs"]["fieldType"], "Field")
        self.assertEqual(payload["style"]["styleID"], "http://www.zotero.org/styles/vancouver")
        self.assertFalse(payload["style"]["bibliographyStyleHasBeenSet"])

    def test_style_urls_become_zotero_style_ids(self) -> None:
        for style in ("apa", "https://www.zotero.org/styles/apa", "http://zotero.org/styles/apa/", "https://citationstyles.org/styles/apa"):
            with self.subTest(style=style):
                self.assertEqual(style_id(style), "http://www.zotero.org/styles/apa")
                self.assertEqual(style_name(style), "apa")
        for style in ("https://example.com/styles/apa", "../apa", ""):
            with self.subTest(style=style), self.assertRaises(ValueError):
                style_id(style)

    def test_docx_contains_plugin_fields(self) -> None:
        source = self.root / "draft.docx"
        document = Document()
        document.add_paragraph("Before {{zotero:ABCD1234|locator=4|label=page}} after.")
        document.add_paragraph("{{zotero:bibliography}}")
        document.core_properties.author = "Local User"
        document.save(source)
        destination = self.root / "draft.zotero.docx"
        stats = embed_document(
            source,
            destination,
            user_id="1",
            style="vancouver",
            locale="en-US",
            items={"ABCD1234": _sample_item()},
            insert_bibliography=True,
        )
        self.assertEqual(stats.citation_count, 1)
        self.assertTrue(stats.bibliography_inserted)
        self.assertEqual(Document(source).paragraphs[0].text, "Before {{zotero:ABCD1234|locator=4|label=page}} after.")
        with zipfile.ZipFile(destination) as package:
            codes = _field_codes(package.read("word/document.xml"))
            preferences = _preferences(package.read("docProps/custom.xml"))
            core = package.read("docProps/core.xml").decode("utf-8")
        self.assertTrue(codes[0].startswith(" ADDIN ZOTERO_ITEM CSL_CITATION "))
        self.assertTrue(codes[0].endswith(" "))
        citation = json.loads(codes[0][codes[0].index("{") : codes[0].rindex("}") + 1])
        self.assertEqual(citation["schema"], "https://github.com/citation-style-language/schema/raw/master/csl-citation.json")
        self.assertEqual(citation["citationItems"][0]["uris"], ["http://zotero.org/users/1/items/ABCD1234"])
        self.assertEqual(citation["citationItems"][0]["locator"], "4")
        self.assertEqual(citation["properties"]["plainCitation"], "1, p. 4")
        self.assertIn("Example & <Trial>", json.dumps(citation))
        self.assertTrue(codes[1].startswith(" ADDIN ZOTERO_BIBL "))
        self.assertTrue(codes[1].rstrip().endswith("CSL_BIBLIOGRAPHY"))
        self.assertEqual(preferences["prefs"]["noteType"], 0)
        self.assertTrue(preferences["style"]["hasBibliography"])
        self.assertNotIn("Local User", core)

    def test_prefix_goes_on_the_first_item_and_suffix_on_the_last(self) -> None:
        job = _job({"AAAA1111": _sample_item(), "BBBB2222": _sample_item("Other", "Jones", 2021)}, numeric=False)
        marker = parse_marker_body("AAAA1111+BBBB2222|prefix=see|locator=4|suffix=, emphasis added").citation
        assert marker is not None
        visible = _visible_citation(job, marker)
        self.assertEqual(visible, "(see Smith, 2020; Jones, 2021, p. 4, emphasis added)")
        instruction = citation_instruction(job, marker, visible)
        first, last = json.loads(instruction[instruction.index("{") : instruction.rindex("}") + 1])["citationItems"]
        self.assertEqual(first["prefix"], "see")
        self.assertFalse({"suffix", "locator"} & set(first))
        self.assertEqual((last["locator"], last["label"], last["suffix"]), ("4", "page", ", emphasis added"))
        self.assertNotIn("prefix", last)

    def test_suppress_author_applies_to_every_item(self) -> None:
        job = _job({"AAAA1111": _sample_item(), "BBBB2222": _sample_item("Other", "Jones", 2021)}, numeric=False)
        marker = parse_marker_body("AAAA1111+BBBB2222|suppress-author=true").citation
        assert marker is not None
        visible = _visible_citation(job, marker)
        self.assertEqual(visible, "(2020; 2021)")
        instruction = citation_instruction(job, marker, visible)
        entries = json.loads(instruction[instruction.index("{") : instruction.rindex("}") + 1])["citationItems"]
        self.assertTrue(all(entry["suppress-author"] for entry in entries))

    def test_numeric_placeholder_bibliography_is_numbered_in_order(self) -> None:
        job = _job(
            {
                "AAAA1111": EmbeddedItem(csl={}, bibliography="1. Smith A. First."),
                "BBBB2222": EmbeddedItem(csl={}, bibliography="[1] B. Jones, Second."),
            },
            numeric=True,
        )
        self.assertEqual(_bibliography_text(job), "1. Smith A. First.\n2. B. Jones, Second.")

    def test_citation_keys_finds_every_marker(self) -> None:
        source = self.root / "draft.docx"
        _write_many_markers(source)
        self.assertEqual(citation_keys(source), _MANY_KEYS)

    def test_embed_after_citation_keys_writes_every_marker(self) -> None:
        source = self.root / "draft.docx"
        destination = self.root / "draft.zotero.docx"
        _write_many_markers(source)
        # embed_zotero_word_fields fetches items only for the keys citation_keys returns.
        keys = citation_keys(source)
        stats = embed_document(
            source,
            destination,
            user_id="1",
            style="vancouver",
            locale="en-US",
            items={key: _sample_item() for key in keys},
            insert_bibliography=False,
        )
        with zipfile.ZipFile(destination) as package:
            codes = _field_codes(package.read("word/document.xml"))
        uris = [json.loads(code[code.index("{") : code.rindex("}") + 1])["citationItems"][0]["uris"][0] for code in codes]
        self.assertEqual(stats.citation_count, 90)
        self.assertEqual(uris, [f"http://zotero.org/users/1/items/{key}" for key in _MANY_KEYS])

    def test_header_markers_are_converted_and_no_parts_are_added(self) -> None:
        plain = self._embed(_write_docx(self.root / "plain.docx", "Body {{zotero:ABCD1234}}."), {"ABCD1234": _sample_item()})
        with zipfile.ZipFile(plain) as package:
            self.assertEqual([name for name in package.namelist() if name.startswith(("word/header", "word/footer"))], [])
        document = Document()
        document.add_paragraph("Body {{zotero:ABCD1234}}.")
        document.sections[0].header.paragraphs[0].text = "Header {{zotero:EFGH5678}}"
        source = self.root / "header.docx"
        document.save(source)
        output = self._embed(source, {"ABCD1234": _sample_item(), "EFGH5678": _sample_item()})
        with zipfile.ZipFile(output) as package:
            headers = [package.read(name) for name in package.namelist() if name.startswith("word/header")]
        self.assertEqual(len(headers), 1)
        self.assertIn("ADDIN ZOTERO_ITEM", "".join(_field_codes(headers[0])))

    def test_existing_custom_properties_are_kept_and_old_preferences_replaced(self) -> None:
        source = _write_docx(self.root / "draft.docx", "Claim {{zotero:ABCD1234}}.")
        word._write_preferences(source, ["old", "stale"])
        custom = (
            f'<Properties xmlns="{_CUSTOM}" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
            f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="2" name="Client"><vt:lpwstr>Acme</vt:lpwstr></property>'
            f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="3" name="ZOTERO_PREF_1"><vt:lpwstr>old</vt:lpwstr></property>'
            f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="4" name="ZOTERO_PREF_2"><vt:lpwstr>stale</vt:lpwstr></property>'
            "</Properties>"
        )
        _replace_part(source, "docProps/custom.xml", custom.encode("utf-8"))
        output = self._embed(source, {"ABCD1234": _sample_item()})
        with zipfile.ZipFile(output) as package:
            root = ElementTree.fromstring(package.read("docProps/custom.xml"))
        properties = {child.attrib["name"]: "".join(child.itertext()) for child in root}
        pids = [child.attrib["pid"] for child in root]
        self.assertEqual(properties["Client"], "Acme")
        self.assertNotIn("stale", properties.values())
        self.assertEqual(_preferences(ElementTree.tostring(root))["dataVersion"], 4)
        self.assertEqual(len(pids), len(set(pids)))

    def test_a_failed_save_leaves_no_temporary_files(self) -> None:
        source = _write_docx(self.root / "draft.docx", "Claim {{zotero:ABCD1234}}.")
        with mock.patch.object(word, "_write_preferences", side_effect=ValueError("broken")), self.assertRaises(ValueError):
            self._embed(source, {"ABCD1234": _sample_item()})
        self.assertEqual(sorted(path.name for path in self.root.iterdir()), ["draft.docx"])

    def test_instruction_is_plugin_shaped(self) -> None:
        job = _job({"ABCD1234": _sample_item("Title")}, numeric=False)
        marker = parse_marker_body("ABCD1234")
        assert marker.citation is not None
        instruction = citation_instruction(job, marker.citation, "(Smith, 2020)")
        self.assertIn(" ADDIN ZOTERO_ITEM CSL_CITATION ", instruction)


class EmbedToolTests(unittest.TestCase):
    """embed_zotero_word_fields, with Zotero replaced by a stub."""

    def setUp(self) -> None:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.source = _write_docx(self.root / "draft.docx", "Claim {{zotero:ABCD1234}}.")
        for patcher in (
            mock.patch.dict(os.environ, {"ZOTERO_USER_ID": "1"}),
            mock.patch.object(client, "_windows_environment", return_value=""),
            mock.patch.object(client.logger, "disabled", True),
        ):
            patcher.start()
            self.addCleanup(patcher.stop)
        fetch = mock.patch("zotero_mcp.server.cited_items", return_value={"ABCD1234": (_sample_item().csl, _sample_item().bibliography)})
        self.fetch = fetch.start()
        self.addCleanup(fetch.stop)

    def test_embed_tool_requires_user_id(self) -> None:
        with mock.patch.dict(os.environ, {"ZOTERO_USER_ID": ""}):
            result = embed_zotero_word_fields(str(self.source))
        self.assertEqual(result.status, "Unconfigured")

    def test_the_tool_writes_fields(self) -> None:
        result = embed_zotero_word_fields(str(self.source), style="https://www.zotero.org/styles/apa")
        self.assertEqual((result.status, result.citation_count), ("OK", 1))
        self.assertEqual(result.style_id, "http://www.zotero.org/styles/apa")
        self.assertTrue((self.root / "draft.zotero.docx").is_file())
        self.fetch.assert_called_once_with(["ABCD1234"], "apa")

    def test_arguments_are_checked_before_zotero_is_read(self) -> None:
        original = self.source.read_bytes()
        cases = {
            "locale": {"locale": "english"},
            "missing folder": {"output_path": str(self.root / "missing" / "out.docx")},
            "not docx": {"output_path": str(self.root / "out.doc")},
            "the source itself": {"output_path": str(self.source)},
            "style": {"style": "https://example.com/apa"},
        }
        for name, arguments in cases.items():
            with self.subTest(name):
                result = embed_zotero_word_fields(str(self.source), **arguments)
                self.assertEqual(result.status, "API_Error")
        self.assertIn("full paths", embed_zotero_word_fields("draft.docx").message)
        self.fetch.assert_not_called()
        self.assertEqual(self.source.read_bytes(), original)

    def test_unreadable_documents_are_explained(self) -> None:
        broken = self.root / "broken.docx"
        broken.write_bytes(b"This is not a Word document.")
        self.assertIn("not a readable Word .docx file", embed_zotero_word_fields(str(broken)).message)
        legacy = self.root / "old.doc"
        legacy.write_bytes(b"\xd0\xcf\x11\xe0")
        self.assertIn("Save a .doc file as .docx", embed_zotero_word_fields(str(legacy)).message)
        self.fetch.assert_not_called()

    def test_missing_items_are_named(self) -> None:
        self.fetch.side_effect = LookupError("No library item uses ABCD1234.")
        result = embed_zotero_word_fields(str(self.source))
        self.assertEqual(result.status, "API_Error")
        self.assertIn("ABCD1234", result.message)

    def test_a_locked_output_is_explained(self) -> None:
        locked = PermissionError(13, "Permission denied", str(self.root / "draft.zotero.docx"))
        with mock.patch("zotero_mcp.server.embed_document", side_effect=locked):
            result = embed_zotero_word_fields(str(self.source))
        self.assertEqual(result.status, "API_Error")
        self.assertIn("close it and try again", result.message)


def _replace_part(path: Path, name: str, data: bytes) -> None:
    with zipfile.ZipFile(path) as package:
        parts = {info.filename: package.read(info.filename) for info in package.infolist()}
    parts[name] = data
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        for part, content in parts.items():
            package.writestr(part, content)


def _field_codes(document_xml: bytes) -> list[str]:
    root = ElementTree.fromstring(document_xml)
    codes = []
    for paragraph in root.iter(f"{{{_W}}}p"):
        texts = [node.text or "" for node in paragraph.iter(f"{{{_W}}}instrText")]
        if texts:
            codes.append("".join(texts))
    return codes


def _preferences(custom_xml: bytes) -> dict:
    root = ElementTree.fromstring(custom_xml)
    chunks = []
    for child in root:
        name = child.attrib.get("name", "")
        if not name.startswith("ZOTERO_PREF_"):
            continue
        chunks.append((int(name.rsplit("_", 1)[-1]), "".join(child.itertext())))
    return json.loads("".join(text for _index, text in sorted(chunks)))


if __name__ == "__main__":
    unittest.main()
