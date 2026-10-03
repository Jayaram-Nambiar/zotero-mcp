"""Write Microsoft Word fields that the Zotero plugin can refresh.

Zotero stores a citation as a Word field whose instruction begins
`` ADDIN ZOTERO_ITEM CSL_CITATION `` followed by a citation in the CSL citation
JSON schema, a bibliography as `` ADDIN ZOTERO_BIBL ... CSL_BIBLIOGRAPHY ``,
and document preferences in custom properties ``ZOTERO_PREF_1``,
``ZOTERO_PREF_2``, and so on. Each property holds at most 255 characters. The
plugin concatenates them and parses one JSON object (dataVersion 4). Zotero's
field-code note introduces these fields:
https://www.zotero.org/support/kb/word_field_codes
"""

from __future__ import annotations

import json
import os
import re
import secrets
import tempfile
import zipfile
import zlib
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
from xml.etree import ElementTree

from docx import Document
from docx.opc.exceptions import OpcError
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tc
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

_MARKER = re.compile(r"\{\{\s*zotero:([^{}]+?)\s*\}\}")
_KEY = re.compile(r"^[A-Za-z0-9]{8}$")
_STYLE_NAME = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_-]*$")
_LOCALE = re.compile(r"^[A-Za-z]{2,3}(?:-[A-Za-z0-9]{2,8})*$")
_LEADING_NUMBER = re.compile(r"^\s*(?:\[\d+\]|\d+\.)\s*")
_OPTION_NAMES = frozenset({"locator", "label", "prefix", "suffix", "suppress-author"})
_LABELS = frozenset({
    "page", "book", "chapter", "column", "figure", "folio", "issue", "line",
    "note", "opus", "paragraph", "part", "section", "sub verbo", "volume", "verse",
})
_NUMERIC_STYLES = frozenset({
    "vancouver",
    "vancouver-superscript",
    "ieee",
    "nature",
    "bmj",
    "american-medical-association",
    "nlm",
    "nlm-brackets",
})
_SIZE_LIMIT = 30 * 1024 * 1024
_PREF_LIMIT = 255
_CITATION_SCHEMA = "https://github.com/citation-style-language/schema/raw/master/csl-citation.json"
_CUSTOM_NS = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
_VT_NS = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
_CUSTOM_TYPE = "application/vnd.openxmlformats-officedocument.custom-properties+xml"
_CUSTOM_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties"
_FMTID = "{D5CDD505-2E9C-101B-9397-08002B2CF9AE}"
# What a damaged or non-Word file raises while it is read.
_UNREADABLE = (OpcError, KeyError, zipfile.BadZipFile, zlib.error, SyntaxError)

ElementTree.register_namespace("", _CUSTOM_NS)
ElementTree.register_namespace("vt", _VT_NS)


@dataclass
class CitationMarker:
    keys: list[str]
    locator: str = ""
    label: str = ""
    prefix: str = ""
    suffix: str = ""
    suppress_author: bool = False


@dataclass
class Marker:
    kind: str
    citation: CitationMarker | None = None


@dataclass
class EmbeddedItem:
    csl: dict[str, Any]
    bibliography: str


@dataclass
class EmbedStats:
    output_path: Path
    style_id: str
    citation_count: int
    bibliography_inserted: bool


@dataclass
class _Job:
    user_id: str
    style_id: str
    locale: str
    items: dict[str, EmbeddedItem]
    order: dict[str, int]
    numeric: bool
    session_id: str
    citation_count: int = 0
    bibliography_inserted: bool = False
    _ids: set[str] = field(default_factory=set)


def style_id(style: str) -> str:
    """Return the CSL style ID Zotero uses, such as http://www.zotero.org/styles/apa."""
    value = style.strip()
    if value.startswith(("http://", "https://")):
        parsed = urlparse(value)
        host = parsed.netloc.lower().removeprefix("www.")
        if host not in {"zotero.org", "citationstyles.org"}:
            raise ValueError("style must be a short name such as vancouver, or a zotero.org or citationstyles.org style URL.")
        value = parsed.path.rstrip("/").rsplit("/", 1)[-1]
    if _STYLE_NAME.fullmatch(value) is None:
        raise ValueError("style must be a short name such as vancouver or apa.")
    return f"http://www.zotero.org/styles/{value}"


def style_name(style: str) -> str:
    return style_id(style).rsplit("/", 1)[-1]


def check_locale(locale: str) -> None:
    if _LOCALE.fullmatch(locale) is None:
        raise ValueError("locale must look like en-US.")


def check_paths(source: Path, destination: Path) -> None:
    """Refuse paths the conversion cannot use, before any work is done."""
    if not source.is_absolute() or not destination.is_absolute():
        raise ValueError("Pass full paths, such as C:\\drafts\\paper.docx or /Users/you/paper.docx.")
    if source.suffix.lower() != ".docx" or destination.suffix.lower() != ".docx":
        raise ValueError("Both paths must be .docx files. Save a .doc file as .docx in Word first.")
    if not source.is_file():
        raise ValueError("docx_path is not a file.")
    if source.stat().st_size > _SIZE_LIMIT:
        raise ValueError("The document is larger than the 30 MB limit for this tool.")
    if not destination.parent.is_dir():
        raise ValueError("The output directory does not exist.")
    if destination.exists() and destination.samefile(source):
        raise ValueError("output_path must differ from docx_path. The original document is never overwritten.")


def _open(source: Path) -> Any:
    try:
        return Document(str(source))
    except _UNREADABLE as exc:
        raise ValueError(f"{source.name} is not a readable Word .docx file.") from exc


def parse_marker_body(body: str) -> Marker:
    text = body.strip()
    if text == "bibliography":
        return Marker(kind="bibliography")
    item_part, separator, option_part = text.partition("|")
    keys = [part.strip() for part in item_part.split("+") if part.strip()]
    if not keys or any(_KEY.fullmatch(key) is None for key in keys):
        raise ValueError(
            "A citation marker needs one or more 8-character Zotero item keys, "
            "for example {{zotero:ABCD1234}} or {{zotero:ABCD1234+EFGH5678}}."
        )
    citation = CitationMarker(keys=keys)
    if separator:
        for piece in option_part.split("|"):
            if "=" not in piece:
                raise ValueError(f"Citation option '{piece}' must look like locator=12.")
            name, value = piece.split("=", 1)
            name = name.strip()
            value = value.strip()
            if name not in _OPTION_NAMES:
                raise ValueError(f"Unknown citation option '{name}'.")
            if name == "locator":
                citation.locator = value
            elif name == "label":
                if value not in _LABELS:
                    raise ValueError(f"Citation label '{value}' is not a CSL locator term.")
                citation.label = value
            elif name == "prefix":
                citation.prefix = value
            elif name == "suffix":
                citation.suffix = value
            elif name == "suppress-author":
                if value.lower() not in {"true", "false"}:
                    raise ValueError("suppress-author must be true or false.")
                citation.suppress_author = value.lower() == "true"
    if citation.locator and not citation.label:
        citation.label = "page"
    return Marker(kind="citation", citation=citation)


def _parse(match: re.Match[str]) -> Marker:
    try:
        return parse_marker_body(match.group(1))
    except ValueError as exc:
        raise ValueError(f"Marker {match.group(0)}: {exc}") from None


def chunk_property(value: str, limit: int = _PREF_LIMIT) -> list[str]:
    """Split document preferences the way the Word plugin stores them."""
    if any(ord(char) < 32 and char not in "\t" for char in value):
        raise ValueError("Document preferences contain a character Word custom properties cannot store.")
    return [value[index:index + limit] for index in range(0, len(value), limit)] or [""]


def document_preferences(*, style: str, locale: str, session_id: str, has_bibliography: bool) -> str:
    check_locale(locale)
    payload = {
        "style": {
            "styleID": style_id(style),
            "locale": locale,
            "hasBibliography": has_bibliography,
            "bibliographyStyleHasBeenSet": False,
        },
        "prefs": {
            "fieldType": "Field",
            "noteType": 0,
        },
        "sessionID": session_id,
        "zoteroVersion": "7.0",
        "dataVersion": 4,
    }
    return json.dumps(payload, separators=(",", ":"))


def citation_instruction(job: _Job, marker: CitationMarker, visible: str) -> str:
    """Build the field instruction. A prefix goes on the first item; locator and suffix on the last."""
    citation_id = _new_id(job)
    items = []
    last = len(marker.keys) - 1
    for index, key in enumerate(marker.keys):
        item_data = dict(job.items[key].csl)
        item_data.setdefault("id", key)
        entry: dict[str, Any] = {
            "id": item_data.get("id", key),
            "uris": [f"http://zotero.org/users/{job.user_id}/items/{key}"],
            "itemData": item_data,
        }
        if index == 0 and marker.prefix:
            entry["prefix"] = marker.prefix
        if index == last:
            if marker.locator:
                entry["locator"] = marker.locator
                entry["label"] = marker.label or "page"
            if marker.suffix:
                entry["suffix"] = marker.suffix
        if marker.suppress_author:
            entry["suppress-author"] = True
        items.append(entry)
    payload = {
        "citationID": citation_id,
        "properties": {
            "formattedCitation": visible,
            "plainCitation": visible,
            "noteIndex": 0,
        },
        "citationItems": items,
        "schema": _CITATION_SCHEMA,
    }
    return f" ADDIN ZOTERO_ITEM CSL_CITATION {json.dumps(payload, separators=(',', ':'))} "


def bibliography_instruction() -> str:
    payload = {"uncited": [], "omitted": [], "custom": []}
    return f" ADDIN ZOTERO_BIBL {json.dumps(payload, separators=(',', ':'))} CSL_BIBLIOGRAPHY "


def _new_id(job: _Job) -> str:
    while True:
        value = secrets.token_hex(4)
        if value not in job._ids:
            job._ids.add(value)
            return value


def _visible_citation(job: _Job, marker: CitationMarker) -> str:
    """A readable stand-in until Zotero's Refresh formats the citation."""
    if job.numeric:
        text = ",".join(str(job.order[key]) for key in marker.keys)
    else:
        parts = []
        for key in marker.keys:
            author, year = _author_year(job.items[key].csl)
            parts.append(year if marker.suppress_author else f"{author}, {year}")
        text = "; ".join(parts)
    if marker.locator:
        label = "p." if marker.label in {"", "page"} else marker.label
        text = f"{text}, {label} {marker.locator}"
    if marker.prefix:
        text = f"{marker.prefix} {text}"
    if marker.suffix:
        text = f"{text}{marker.suffix}"
    return text if job.numeric else f"({text})"


def _author_year(csl: dict[str, Any]) -> tuple[str, str]:
    authors = csl.get("author") or []
    name = "Unknown"
    if authors and isinstance(authors[0], dict):
        person = authors[0]
        name = str(person.get("family") or person.get("literal") or person.get("name") or "Unknown")
    issued = csl.get("issued") if isinstance(csl.get("issued"), dict) else {}
    parts = issued.get("date-parts") if isinstance(issued, dict) else None
    year = "n.d."
    if isinstance(parts, list) and parts and isinstance(parts[0], list) and parts[0]:
        year = str(parts[0][0])
    return name, year


def _bibliography_text(job: _Job) -> str:
    lines = []
    for key, number in sorted(job.order.items(), key=lambda item: item[1]):
        line = job.items[key].bibliography.strip()
        if job.numeric:
            # Zotero numbers a one-item bibliography 1, so renumber each line.
            line = f"{number}. {_LEADING_NUMBER.sub('', line)}"
        lines.append(line)
    return "\n".join(lines)


def iter_paragraphs(document: Any):
    """Yield body, table, header, and footer paragraphs once each."""
    # Hold the elements, not their id()s: lxml frees a proxy that nothing
    # references, and a later proxy for a different element can reuse its id.
    seen: set[CT_P] = set()

    def take(paragraph: Paragraph):
        if paragraph._p in seen:
            return None
        seen.add(paragraph._p)
        return paragraph

    def walk_table(table: Table):
        seen_cells: set[CT_Tc] = set()
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen_cells:
                    continue
                seen_cells.add(cell._tc)
                for paragraph in cell.paragraphs:
                    found = take(paragraph)
                    if found is not None:
                        yield found
                for nested in cell.tables:
                    yield from walk_table(nested)

    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            found = take(Paragraph(child, document))
            if found is not None:
                yield found
        elif child.tag == qn("w:tbl"):
            yield from walk_table(Table(child, document))
    for section in document.sections:
        for container in (section.header, section.footer):
            # A linked header or footer has no part of its own: it repeats the
            # previous section's, or the first section has none. Reading its
            # paragraphs would add an empty part to the document.
            if container.is_linked_to_previous:
                continue
            for paragraph in container.paragraphs:
                found = take(paragraph)
                if found is not None:
                    yield found
            for table in container.tables:
                yield from walk_table(table)


def _markers_in(paragraph: Paragraph) -> list[Marker]:
    return [_parse(match) for match in _MARKER.finditer(paragraph.text or "")]


def _add_field(paragraph: Paragraph, instruction: str, result: str) -> None:
    _field_char(paragraph, "begin")
    run = paragraph.add_run()._r
    for index in range(0, len(instruction), 200):
        node = OxmlElement("w:instrText")
        node.set(qn("xml:space"), "preserve")
        node.text = instruction[index:index + 200]
        run.append(node)
    _field_char(paragraph, "separate")
    _result_run(paragraph, result)
    _field_char(paragraph, "end")


def _field_char(paragraph: Paragraph, kind: str) -> None:
    node = OxmlElement("w:fldChar")
    node.set(qn("w:fldCharType"), kind)
    paragraph.add_run()._r.append(node)


def _result_run(paragraph: Paragraph, result: str) -> None:
    run = paragraph.add_run()
    lines = result.split("\n")
    run.add_text(lines[0] if lines else "")
    for line in lines[1:]:
        run.add_break()
        run.add_text(line)


def _clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _write_paragraph(paragraph: Paragraph, job: _Job) -> None:
    text = paragraph.text or ""
    if _MARKER.search(text) is None:
        return
    pieces: list[tuple[str, Marker | None]] = []
    cursor = 0
    for match in _MARKER.finditer(text):
        if match.start() > cursor:
            pieces.append((text[cursor:match.start()], None))
        pieces.append(("", _parse(match)))
        cursor = match.end()
    if cursor < len(text):
        pieces.append((text[cursor:], None))
    _clear_paragraph(paragraph)
    for plain, marker in pieces:
        if marker is None:
            if plain:
                paragraph.add_run(plain)
            continue
        if marker.kind == "bibliography":
            _add_field(paragraph, bibliography_instruction(), _bibliography_text(job))
            job.bibliography_inserted = True
            continue
        assert marker.citation is not None
        visible = _visible_citation(job, marker.citation)
        _add_field(paragraph, citation_instruction(job, marker.citation, visible), visible)
        job.citation_count += 1


def _scan(paragraphs: list[Paragraph]) -> dict[str, int]:
    """Number item keys by first appearance, and allow one bibliography marker."""
    order: dict[str, int] = {}
    bibliographies = 0
    for paragraph in paragraphs:
        for marker in _markers_in(paragraph):
            if marker.kind == "bibliography":
                bibliographies += 1
                continue
            assert marker.citation is not None
            for key in marker.citation.keys:
                if key not in order:
                    order[key] = len(order) + 1
    if bibliographies > 1:
        raise ValueError("The document has more than one {{zotero:bibliography}} marker.")
    if not order:
        raise ValueError("The document has no {{zotero:ITEMKEY}} citation markers.")
    return order


def citation_keys(source: Path) -> list[str]:
    """Return item keys in first-seen order from citation markers."""
    if not source.is_file():
        raise ValueError("docx_path is not a file.")
    return list(_scan(list(iter_paragraphs(_open(source)))))


def embed_document(
    source: Path,
    destination: Path,
    *,
    user_id: str,
    style: str,
    locale: str,
    items: dict[str, EmbeddedItem],
    insert_bibliography: bool,
) -> EmbedStats:
    """Replace ``{{zotero:...}}`` markers with Zotero Word fields."""
    if not user_id.isascii() or not user_id.isdigit():
        raise ValueError("user_id must be the numeric Zotero user ID.")
    check_paths(source, destination)
    check_locale(locale)
    resolved_style = style_id(style)
    document = _open(source)
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    paragraphs = list(iter_paragraphs(document))
    order = _scan(paragraphs)
    missing = [key for key in order if key not in items]
    if missing:
        raise ValueError("No Zotero item was loaded for " + ", ".join(missing) + ".")
    job = _Job(
        user_id=user_id,
        style_id=resolved_style,
        locale=locale,
        items=items,
        order=order,
        numeric=style_name(style) in _NUMERIC_STYLES,
        session_id=secrets.token_hex(8),
    )
    for paragraph in paragraphs:
        _write_paragraph(paragraph, job)
    if insert_bibliography and not job.bibliography_inserted:
        paragraph = document.add_paragraph()
        _add_field(paragraph, bibliography_instruction(), _bibliography_text(job))
        job.bibliography_inserted = True
    preferences = document_preferences(
        style=style,
        locale=locale,
        session_id=job.session_id,
        has_bibliography=job.bibliography_inserted,
    )
    # A unique temporary name: a retried call may run while a timed-out one is still writing.
    handle, name = tempfile.mkstemp(prefix=f"{destination.stem}.", suffix=".partial", dir=destination.parent)
    os.close(handle)
    temporary = Path(name)
    try:
        document.save(str(temporary))
        _write_preferences(temporary, chunk_property(preferences))
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)
    return EmbedStats(
        output_path=destination,
        style_id=resolved_style,
        citation_count=job.citation_count,
        bibliography_inserted=job.bibliography_inserted,
    )


def _write_preferences(path: Path, chunks: list[str]) -> None:
    handle, name = tempfile.mkstemp(suffix=".zip", dir=path.parent)
    os.close(handle)
    temporary = Path(name)
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w") as target:
            names = set(source.namelist())
            if "[Content_Types].xml" not in names or "_rels/.rels" not in names:
                raise ValueError("The .docx package is missing its content types or relationships.")
            existing = source.read("docProps/custom.xml") if "docProps/custom.xml" in names else None
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "[Content_Types].xml":
                    data = _content_types(data)
                elif info.filename == "_rels/.rels":
                    data = _relationships(data)
                elif info.filename == "docProps/custom.xml":
                    continue
                target.writestr(info, data)
            target.writestr("docProps/custom.xml", _custom_xml(chunks, existing))
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def _content_types(data: bytes) -> bytes:
    text = data.decode("utf-8")
    if 'PartName="/docProps/custom.xml"' in text:
        return data
    override = (
        '<Override PartName="/docProps/custom.xml" '
        f'ContentType="{_CUSTOM_TYPE}"/>'
    )
    if "</Types>" not in text:
        raise ValueError("The .docx content types part is not readable.")
    return text.replace("</Types>", override + "</Types>", 1).encode("utf-8")


def _relationships(data: bytes) -> bytes:
    text = data.decode("utf-8")
    if _CUSTOM_REL in text:
        return data
    relationship = (
        '<Relationship Id="rIdZoteroPrefs" '
        f'Type="{_CUSTOM_REL}" Target="docProps/custom.xml"/>'
    )
    if "</Relationships>" not in text:
        raise ValueError("The .docx relationships part is not readable.")
    return text.replace("</Relationships>", relationship + "</Relationships>", 1).encode("utf-8")


def _custom_xml(chunks: list[str], existing: bytes | None) -> bytes:
    if existing:
        try:
            root = ElementTree.fromstring(existing)
        except ElementTree.ParseError as exc:
            raise ValueError("The document's custom properties are not readable.") from exc
        for child in list(root):
            if child.attrib.get("name", "").startswith("ZOTERO_PREF_"):
                root.remove(child)
    else:
        root = ElementTree.Element(f"{{{_CUSTOM_NS}}}Properties")
    used = {int(child.attrib["pid"]) for child in root if child.attrib.get("pid", "").isdigit()}
    next_pid = max(used, default=1) + 1
    for index, chunk in enumerate(chunks, start=1):
        while next_pid in used:
            next_pid += 1
        prop = ElementTree.SubElement(
            root,
            f"{{{_CUSTOM_NS}}}property",
            {"fmtid": _FMTID, "pid": str(next_pid), "name": f"ZOTERO_PREF_{index}"},
        )
        ElementTree.SubElement(prop, f"{{{_VT_NS}}}lpwstr").text = chunk
        used.add(next_pid)
        next_pid += 1
    return ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
